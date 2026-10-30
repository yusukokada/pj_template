"""データ定義書（Excel）の読み書きと点検。

データ定義書の正は docs/ の Excel。AI が直すときは、Excel を JSON に書き出し（dump）、JSON を直し、
新しい版の Excel を作る（build）。Excel の体裁はこのスクリプトが決めるので、毎回同じ形になる。

  python doc_tools/definition.py dump  <定義書.xlsx> <出力.json>
  python doc_tools/definition.py build <入力.json> <新しい版.xlsx>      （出力先が既にあれば止める）
  python doc_tools/definition.py check <定義書.xlsx または .json>
  python doc_tools/definition.py new   <出力.json>                      （空の雛形の JSON）

JSON の形は doc_tools/examples/sample_definition.json を見る。"""
import json
import keyword
import re
import sys
from pathlib import Path

TYPES = {
    "ID": "行を識別する値（例：P1-0001）",
    "コード": "短い記号（例：RUN）",
    "文字": "名称・説明などの文字",
    "区分": "決まった値から選ぶもの。選択肢は 91_コード値案 のコード種別",
    "数値": "工数・件数などの数",
    "比率": "0〜1 の割合",
    "金額": "円単位の整数",
    "日付": "年月日",
    "日時": "年月日と時刻",
    "年月": "YYYY-MM",
    "フラグ": "はい／いいえ",
}
LAYERS = {
    "マスタ": "ほかの記録から参照される、変化の少ないデータ（組織・担当者など）",
    "記録": "業務で発生した事実（タスク・工数・支払いなど）。人やシステムが入力する",
    "導出": "記録とマスタから計算で作るデータ（集計・指標）。人は入力しない",
}
SOURCES = ("入力", "取込", "計算")
# アプリ名に使えない名前（Python・Django の標準モジュールと衝突する）
RESERVED_APPS = {"ast", "io", "os", "re", "json", "test", "tests", "types", "code", "admin", "auth", "core", "work",
                 "config", "tools", "static", "templates", "site", "string", "select", "time", "calendar", "email"}

TABLE_COLS = [("ID", "id", 8), ("領域", "area", 16), ("層", "layer", 8), ("テーブル名", "name", 22),
              ("管理するデータ", "data", 48), ("用途", "use", 48), ("1行の単位", "unit", 22), ("正とする所", "home", 20)]
FIELD_COLS = [("テーブルID", "table", 10), ("テーブル名", "table_name", 18), ("No", "no", 5), ("項目名", "name", 22),
              ("型", "type", 8), ("キー", "key", 6), ("参照先", "ref", 18), ("必須", "required", 6), ("取得", "source", 8),
              ("コード種別", "code", 16), ("説明", "desc", 60)]
AREA_COLS = [("領域", "name", 20), ("アプリ名", "app", 12), ("管理画面での表示名", "label", 30)]
HEAD_ROW = 3   # 1行目に表題、3行目に見出し、4行目から本文


# ---------------------------------------------------------------- 読む
def load(path):
    """Excel か JSON を読み、同じ形の dict にして返す。"""
    path = Path(path)
    if path.suffix.lower() == ".json":
        return json.loads(path.read_text(encoding="utf-8"))
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def rows(sheet, cols):
        out = []
        it = wb[sheet].iter_rows(min_row=HEAD_ROW, values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it)]
        idx = {h: i for i, h in enumerate(head)}
        for r in it:
            if r is None or all(v is None for v in r):
                continue
            out.append({k: (r[idx[h]] if h in idx and idx[h] < len(r) else None) for h, k, _ in cols})
        return out

    legend = {r[0]: r[1] for r in wb["00_凡例"].iter_rows(min_row=HEAD_ROW + 1, values_only=True) if r and r[0]}
    areas = rows("02_領域", AREA_COLS)
    tables = rows("01_テーブル一覧", TABLE_COLS)
    by_id = {t["id"]: dict(t, fields=[]) for t in tables if t["id"]}
    for f in rows("90_全項目", FIELD_COLS):
        if not f["table"]:
            continue
        ref = str(f["ref"]).split()[0] if f["ref"] else ""
        by_id[f["table"]]["fields"].append({
            "name": str(f["name"]), "type": f["type"], "key": f["key"] or "", "ref": ref,
            "required": f["required"] == "○", "source": f["source"] or "", "code": f["code"] or "",
            "desc": f["desc"] or ""})
    codes, kind = {}, None
    for r in wb["91_コード値案"].iter_rows(min_row=HEAD_ROW + 1, values_only=True):
        if not r or r[2] is None:
            continue
        kind = r[0] or kind
        codes.setdefault(kind, []).append(str(r[2]))
    for t in by_id.values():
        for k in ("data", "use", "unit", "home"):
            t[k] = t[k] or ""
    return {"title": legend.get("名称", ""), "version": str(legend.get("版", "")),
            "areas": [{k: a[k] or "" for k in ("name", "app", "label")} for a in areas],
            "tables": list(by_id.values()), "codes": codes}


# ---------------------------------------------------------------- 点検
def check(d):
    """誤りの一覧と、気づき（誤りではないが確かめた方がよい点）の一覧を返す。"""
    err, warn = [], []
    areas = {a["name"]: a for a in d["areas"]}
    apps = [a["app"] for a in d["areas"]]
    for a in d["areas"]:
        if not re.fullmatch(r"[a-z][a-z0-9_]*", a["app"] or ""):
            err.append(f"02_領域「{a['name']}」のアプリ名「{a['app']}」は英小文字・数字・_ で書く")
        elif a["app"] in RESERVED_APPS or keyword.iskeyword(a["app"]):
            err.append(f"02_領域のアプリ名「{a['app']}」は Python・Django の名前と衝突する。別の名前にする")
    if len(set(apps)) != len(apps):
        err.append("02_領域のアプリ名が重複している")
    ids = [t["id"] for t in d["tables"]]
    for x in {i for i in ids if ids.count(i) > 1}:
        err.append(f"テーブルID {x} が重複している")
    T = {t["id"]: t for t in d["tables"]}
    for t in d["tables"]:
        where = f"{t['id']} {t['name']}"
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9]*", str(t["id"])):
            err.append(f"{where}：テーブルID は英字で始まる英数字にする")
        if t["area"] not in areas:
            err.append(f"{where}：領域「{t['area']}」が 02_領域 にない")
        if t["layer"] not in LAYERS:
            err.append(f"{where}：層「{t['layer']}」は {'・'.join(LAYERS)} のどれか")
        for k, label in (("unit", "1行の単位"), ("home", "正とする所"), ("data", "管理するデータ"), ("use", "用途")):
            if not t.get(k):
                warn.append(f"{where}：{label}が空")
        if not t["fields"]:
            err.append(f"{where}：項目がない")
            continue
        if not any("PK" in f["key"] for f in t["fields"]):
            err.append(f"{where}：キー（PK）の項目がない")
        names = [f["name"] for f in t["fields"]]
        for x in {n for n in names if names.count(n) > 1}:
            err.append(f"{where}：項目名「{x}」が重複している")
        for f in t["fields"]:
            fw = f"{where}.{f['name']}"
            if f["type"] not in TYPES:
                err.append(f"{fw}：型「{f['type']}」は {'・'.join(TYPES)} のどれか")
            if f["ref"] and f["ref"] not in T:
                err.append(f"{fw}：参照先 {f['ref']} がない")
            if f["source"] not in SOURCES:
                warn.append(f"{fw}：取得「{f['source']}」は {'・'.join(SOURCES)} のどれか")
            if t["layer"] == "導出" and f["source"] == "入力" and "PK" not in f["key"]:
                warn.append(f"{fw}：導出のテーブルに人が入力する項目がある")
            if t["layer"] != "導出" and f["source"] == "計算":
                warn.append(f"{fw}：計算で出す値を記録・マスタに持っている（導出のテーブルに分けるか確かめる）")
            if f["type"] == "区分":
                kind = f["code"] or f["name"]
                if kind not in d["codes"]:
                    warn.append(f"{fw}：コード種別「{kind}」が 91_コード値案 にない（自由入力になる）")
            elif f["code"]:
                warn.append(f"{fw}：区分でない項目にコード種別がある")
    for kind, vals in d["codes"].items():
        for v in vals:
            if "：" not in v:
                warn.append(f"コード値案「{kind}」の「{v}」は「コード：表示名」の形にする")
    used = {f["code"] or f["name"] for t in d["tables"] for f in t["fields"] if f["type"] == "区分"}
    for kind in d["codes"]:
        if kind not in used:
            warn.append(f"コード値案「{kind}」はどの項目からも使われていない")
    return err, warn


# ---------------------------------------------------------------- 書く
def build(d, out):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    out = Path(out)
    assert not out.exists(), f"出力先が既にある：{out}"
    err, _ = check(d)
    assert not err, "定義に誤りがあるため作らない：\n" + "\n".join(err)
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="75787B")
    head_font = Font(bold=True, color="FFFFFF", name="Meiryo UI", size=10)
    body_font = Font(name="Meiryo UI", size=10)
    thin = Side(style="thin", color="D0D0CE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(name, title, cols, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, name="Meiryo UI", size=12)
        for j, (h, _, w) in enumerate(cols, 1):
            c = ws.cell(HEAD_ROW, j, h)
            c.fill, c.font, c.border = head_fill, head_font, border
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
        for i, r in enumerate(rows, HEAD_ROW + 1):
            for j, v in enumerate(r, 1):
                c = ws.cell(i, j, v)
                c.font, c.border, c.alignment = body_font, border, wrap
        ws.freeze_panes = ws.cell(HEAD_ROW + 1, 1)
        if rows:
            ws.auto_filter.ref = f"A{HEAD_ROW}:{openpyxl.utils.get_column_letter(len(cols))}{HEAD_ROW + len(rows)}"
        return ws

    T = {t["id"]: t for t in d["tables"]}
    legend = [("名称", d["title"]), ("版", d["version"]), ("", ""), ("■ 層", "")]
    legend += [(k, v) for k, v in LAYERS.items()] + [("", ""), ("■ 型", "")] + [(k, v) for k, v in TYPES.items()]
    legend += [("", ""), ("■ 取得", ""), ("入力", "人が画面・Excel で入れる"), ("取込", "ほかのシステムから取り込む"),
               ("計算", "ほかの項目から計算で出す")]
    legend += [("", ""), ("■ キー", "PK＝行を識別する項目（複数あれば組み合わせで識別）。参照先があればほかのテーブルを指す")]
    sheet("00_凡例", "凡例", [("項目", "", 18), ("内容", "", 90)], legend, first=True)
    sheet("01_テーブル一覧", "テーブル一覧", TABLE_COLS + [("項目数", "", 8)],
          [[t[k] for _, k, _ in TABLE_COLS] + [len(t["fields"])] for t in d["tables"]])
    sheet("02_領域", "領域とアプリ名（試作アプリでの分け方）", AREA_COLS, [[a[k] for _, k, _ in AREA_COLS] for a in d["areas"]])
    rows = []
    for t in d["tables"]:
        for n, f in enumerate(t["fields"], 1):
            ref = f"{f['ref']} {T[f['ref']]['name']}" if f["ref"] else ""
            rows.append([t["id"], t["name"], n, f["name"], f["type"], f["key"], ref, "○" if f["required"] else "",
                         f["source"], f["code"], f["desc"]])
    ws = sheet("90_全項目", "全項目", FIELD_COLS, rows)
    n = HEAD_ROW + max(len(rows), 1) + 200
    for col, vals in (("E", list(TYPES)), ("F", ["PK"]), ("H", ["○"]), ("I", list(SOURCES))):
        dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
        dv.add(f"{col}{HEAD_ROW + 1}:{col}{n}")
        ws.add_data_validation(dv)
    crow = []
    for kind, vals in d["codes"].items():
        for i, v in enumerate(vals, 1):
            crow.append([kind if i == 1 else None, i, v])
    sheet("91_コード値案", "コード値案（区分の項目の選択肢。「コード：表示名」で書く）",
          [("コード種別", "", 22), ("No", "", 5), ("コード値（案）", "", 40)], crow)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def empty():
    return {"title": "（プロジェクト名） データ定義書", "version": "0.1",
            "areas": [{"name": "共通マスタ", "app": "master", "label": "00 共通マスタ"}],
            "tables": [], "codes": {}}


def main(argv):
    if len(argv) < 2:
        raise SystemExit(__doc__)
    cmd = argv[0]
    if cmd == "dump":
        d = load(argv[1])
        Path(argv[2]).parent.mkdir(parents=True, exist_ok=True)
        Path(argv[2]).write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{len(d['tables'])} テーブル・{sum(len(t['fields']) for t in d['tables'])} 項目を書き出した：{argv[2]}")
    elif cmd == "build":
        out = build(load(argv[1]), argv[2])
        print(f"作った：{out}")
    elif cmd == "check":
        d = load(argv[1])
        err, warn = check(d)
        for e in err:
            print("NG ", e)
        for w in warn:
            print("気づき", w)
        print(f"{len(d['tables'])} テーブル・{sum(len(t['fields']) for t in d['tables'])} 項目。誤り {len(err)} 件・気づき {len(warn)} 件")
        sys.exit(1 if err else 0)
    elif cmd == "new":
        p = Path(argv[1])
        assert not p.exists(), f"出力先が既にある：{p}"
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(empty(), ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"作った：{p}")
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
